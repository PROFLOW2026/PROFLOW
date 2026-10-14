"""Inspect ProjectFlow_Financial_2026.xlsx structure (read-only)."""
import json
import re
from pathlib import Path
from openpyxl import load_workbook

VOID_IDS = {
    "5642f1de-b5f7-474d-888a-8686b88b2167",
    "7d127420-d37a-469a-97b5-702389de4fee",
    "2807943d-8fcb-41e5-989d-3ee46982f8b0",
    "c16d5480-8422-43b4-9a24-68bc947e8314",
    "e67a342d-bf56-4a79-9b2f-6a77689cf9a5",
    "b59035cf-fe06-49a6-9cb3-97459bed9b4f",
    "989cc273-cd8e-4932-b205-6869d5154b02",
}

p = Path(r"c:\Users\ERAN YOSEF\Desktop\final projects\FINAL-WEB\projectflow\ProjectFlow_Financial_2026.xlsx")
wb = load_workbook(p, data_only=False)
uuid_re = re.compile(
    r"[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}",
    re.I,
)

out = {"sheets": [], "void_hits": [], "expense_id_rows": 0}
for ws in wb.worksheets:
    sheet_info = {
        "title": ws.title,
        "max_row": ws.max_row,
        "max_col": ws.max_column,
        "sheet_view_rtl": getattr(ws.sheet_view, "rightToLeft", None),
    }
    void_in_sheet = []
    expense_rows = 0
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        texts = [str(c.value) if c.value is not None else "" for c in row]
        joined = " ".join(texts)
        for vid in VOID_IDS:
            if vid in joined or vid.replace("-", "") in joined.replace("-", ""):
                void_in_sheet.append({"void_id": vid, "row": row[0].row, "sample": texts[:8]})
        if uuid_re.search(joined):
            expense_rows += 1
    sheet_info["rows_with_uuid"] = expense_rows
    out["sheets"].append(sheet_info)
    out["void_hits"].extend([{**v, "sheet": ws.title} for v in void_in_sheet])

out_path = Path(
    r"c:\Users\ERAN YOSEF\Desktop\final projects\FINAL-WEB\projectflow\.tools\excel-structure.json"
)
out_path.write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps({"sheets": len(out["sheets"]), "void_hits": len(out["void_hits"])}, indent=2))
