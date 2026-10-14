import json
from pathlib import Path
from openpyxl import load_workbook

p = Path(r"c:\Users\ERAN YOSEF\Desktop\final projects\FINAL-WEB\projectflow\ProjectFlow_Financial_2026.xlsx")
out = Path(r"c:\Users\ERAN YOSEF\Desktop\final projects\FINAL-WEB\projectflow\.tools\excel-layout-dump.json")
wb = load_workbook(p, data_only=False)
data = {}
for name in wb.sheetnames:
    ws = wb[name]
    rows = []
    for r in range(1, min(ws.max_row + 1, 50)):
        vals = []
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(r, c)
            v = cell.value
            if v is not None:
                vals.append({"c": c, "v": str(v)[:120]})
        if vals:
            rows.append({"r": r, "cells": vals})
    data[name] = rows
out.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
