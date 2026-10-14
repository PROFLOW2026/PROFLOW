import json
import re
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(r"c:\Users\ERAN YOSEF\Desktop\final projects\FINAL-WEB\projectflow")
XLSX = ROOT / "ProjectFlow_Financial_2026.xlsx"
UUID_RE = re.compile(r"^[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}$", re.I)

wb = load_workbook(XLSX, data_only=True)
out = {"reversals_in_excel": [], "installment_like": [], "negative_net": []}

for ws in wb.worksheets:
    if not ws.title[:2].isdigit():
        continue
    for row in range(1, ws.max_row + 1):
        if ws.cell(row, 2).value != "הוצאה":
            continue
        eid = ws.cell(row, 13).value
        if not eid or not UUID_RE.match(str(eid).strip()):
            continue
        net = ws.cell(row, 5).value
        supplier = ws.cell(row, 3).value
        desc = ws.cell(row, 4).value
        paid = ws.cell(row, 8).value
        bal = ws.cell(row, 9).value
        rec = ws.cell(row, 12).value
        status = ws.cell(row, 11).value
        entry = {
            "sheet": ws.title,
            "row": row,
            "id": str(eid).lower(),
            "supplier": supplier,
            "desc": desc,
            "net": net,
            "paid": paid,
            "balance": bal,
            "recognized": rec,
            "status": status,
        }
        if net is not None and float(net) < 0:
            out["negative_net"].append(entry)
        if supplier and "ארכה" in str(supplier) and net and abs(float(net) - 2097.46) < 1:
            out["reversals_in_excel"].append(entry)
        if paid and net and float(paid) > 0 and float(paid) < float(net) * 1.2:
            if bal and float(bal) > 0:
                out["installment_like"].append(entry)

Path(ROOT / ".tools/excel-scan-current.json").write_text(
    json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8"
)
print(json.dumps({k: len(v) for k, v in out.items()}, indent=2))
