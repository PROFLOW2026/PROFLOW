import json
import re
from pathlib import Path
from openpyxl import load_workbook

XLSX = Path(r"c:\Users\ERAN YOSEF\Desktop\final projects\FINAL-WEB\projectflow\ProjectFlow_Financial_2026.xlsx")
UUID_RE = re.compile(r"^[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}$", re.I)
wb = load_workbook(XLSX, data_only=True)
rows = []
for ws in wb.worksheets:
    if not ws.title[:2].isdigit():
        continue
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 2).value != "הוצאה":
            continue
        eid = ws.cell(r, 13).value
        if not eid or not UUID_RE.match(str(eid).strip()):
            continue
        rows.append({
            "sheet": ws.title,
            "row": r,
            "id": str(eid).lower(),
            "supplier": ws.cell(r, 3).value,
            "net": ws.cell(r, 5).value,
            "recognized": ws.cell(r, 12).value,
            "paid": ws.cell(r, 8).value,
            "balance": ws.cell(r, 9).value,
        })
Path(r"c:\Users\ERAN YOSEF\Desktop\final projects\FINAL-WEB\projectflow\.tools\excel-all-expense-rows.json").write_text(
    json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8"
)
print(len(rows))
