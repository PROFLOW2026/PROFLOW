"""
Populate section 0 (actual payments) on monthly Hebrew sheets from org-all-payments-2026.json.
Opens workbook in Microsoft Excel after save.
"""
from __future__ import annotations

import json
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(r"c:\Users\ERAN YOSEF\Desktop\final projects\FINAL-WEB\projectflow")
XLSX = ROOT / "ProjectFlow_Financial_2026.xlsx"
DATA = ROOT / ".tools" / "org-all-payments-2026.json"
REPORT = ROOT / ".tools" / "excel-full-payments-closeout-report.json"

SHEETS = {
    "2026-01": "01 ינואר",
    "2026-02": "02 פברואר",
    "2026-03": "03 מרץ",
    "2026-04": "04 אפריל",
    "2026-05": "05 מאי",
    "2026-06": "06 יוני",
    "2026-07": "07 יולי",
    "2026-08": "08 אוגוסט",
    "2026-09": "09 ספטמבר",
    "2026-10": "10 אוקטובר",
}

HEADERS = [
    "תאריך תשלום",
    "סוג",
    "ספק / עובד",
    "תיאור",
    "NET ₪",
    "מע״מ ₪",
    "GROSS ₪",
    "שיוך / פרויקט",
    "מזהה מקור",
    "הערה",
]

KIND_HE = {
    "expense": "הוצאה — תשלום",
    "payroll": "שכר — מזומן",
    "customer_receipt": "תקבול לקוח",
    "expense_date_unknown": "חריג — תאריך לא מאומת",
}


def q2(v: float) -> float:
    return float(Decimal(str(v)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def find_section1_row(ws) -> int:
    for r in range(1, min(ws.max_row + 1, 200)):
        v = ws.cell(r, 1).value
        if v and str(v).startswith("1 |"):
            return r
    return 4


def delete_old_section0(ws, section1_row: int) -> int:
    start = None
    for r in range(1, section1_row):
        v = ws.cell(r, 1).value
        if v and str(v).startswith("0 |"):
            start = r
            break
    if start is None:
        return 0
    end = section1_row - 1
    count = end - start + 1
    ws.delete_rows(start, count)
    return count


def main() -> None:
    lock = ROOT / "~$ProjectFlow_Financial_2026.xlsx"
    if lock.exists():
        raise SystemExit("Close Excel lock file and retry.")

    payload = json.loads(DATA.read_text(encoding="utf-8"))
    by_month = payload["byMonth"]
    date_unknown = payload.get("dateUnknown", [])
    commitments = payload.get("commitments", [])
    emc = payload.get("emcByMonth", {})

    wb = load_workbook(XLSX)
    stats = []

    for ym, sheet_name in SHEETS.items():
        ws = wb[sheet_name]
        ws.sheet_view.rightToLeft = True
        s1 = find_section1_row(ws)
        deleted = delete_old_section0(ws, s1)
        s1 = find_section1_row(ws)
        all_rows = by_month.get(ym, [])
        out_rows = [p for p in all_rows if p.get("kind") in ("expense", "payroll")]
        in_rows = [p for p in all_rows if p.get("kind") == "customer_receipt"]
        truth = (payload.get("monthlyTruth") or {}).get(ym) or {}
        block = 2 + 1 + len(out_rows) + 1 + (2 + 1 + len(in_rows) if in_rows else 0) + 2
        ws.insert_rows(s1, block)
        r = s1
        ws.cell(r, 1).value = (
            f"0 | תשלומים יוצאים (הוצאות + שכר) — {len(out_rows)} רשומות"
        )
        r += 1
        for c, h in enumerate(HEADERS, start=1):
            ws.cell(r, c).value = h
        r += 1
        for pay in sorted(out_rows, key=lambda x: x.get("paymentDate") or ""):
            ws.cell(r, 1).value = pay.get("paymentDate")
            ws.cell(r, 2).value = KIND_HE.get(pay.get("kind"), pay.get("kind"))
            ws.cell(r, 3).value = pay.get("counterparty")
            ws.cell(r, 4).value = pay.get("description")
            ws.cell(r, 5).value = q2(pay.get("net", 0))
            ws.cell(r, 6).value = q2(pay.get("vat", 0))
            ws.cell(r, 7).value = q2(pay.get("gross", 0))
            ws.cell(r, 8).value = pay.get("allocation")
            ws.cell(r, 9).value = pay.get("sourceId")
            ws.cell(r, 10).value = pay.get("note")
            r += 1
        exp = truth.get("expensesOut") or {}
        pay = truth.get("payrollOut") or {}
        ws.cell(r, 1).value = (
            f"סיכום יוצאים: הוצאות GROSS {q2(exp.get('gross', 0)):,.2f} | שכר GROSS {q2(pay.get('gross', 0)):,.2f} | "
            f"סה״כ NET {q2(truth.get('outgoingPaymentsNet', 0)):,.2f} | מע״מ {q2(truth.get('outgoingPaymentsVat', 0)):,.2f} | "
            f"GROSS {q2(truth.get('outgoingPaymentsGross', 0)):,.2f}"
        )
        r += 1
        if in_rows:
            ws.cell(r, 1).value = f"0a | תקבולים מלקוחות — {len(in_rows)} רשומות (לא חלק מסך הוצאות)"
            r += 1
            for c, h in enumerate(HEADERS, start=1):
                ws.cell(r, c).value = h
            r += 1
            rec = truth.get("receiptsIn") or {}
            sum_in = 0.0
            for pay in sorted(in_rows, key=lambda x: x.get("paymentDate") or ""):
                ws.cell(r, 1).value = pay.get("paymentDate")
                ws.cell(r, 2).value = KIND_HE.get("customer_receipt")
                ws.cell(r, 3).value = pay.get("counterparty")
                ws.cell(r, 4).value = pay.get("description")
                ws.cell(r, 5).value = q2(pay.get("net", 0))
                ws.cell(r, 6).value = q2(pay.get("vat", 0))
                ws.cell(r, 7).value = q2(pay.get("gross", 0))
                ws.cell(r, 8).value = pay.get("allocation")
                ws.cell(r, 9).value = pay.get("sourceId")
                ws.cell(r, 10).value = pay.get("note")
                sum_in += pay.get("gross", 0)
                r += 1
            ws.cell(r, 1).value = f"סיכום תקבולים: GROSS {q2(rec.get('gross', sum_in)):,.2f}"
            r += 1
        emc_val = emc.get(ym)
        if emc_val is not None:
            ws.cell(r, 1).value = (
                f"0b | עלות מעסיק (EMC) בחודש — {q2(float(emc_val)):,.2f} ₪ (נפרד מתשלומי שכר)"
            )
        stats.append(
            {
                "ym": ym,
                "outgoing": len(out_rows),
                "receipts": len(in_rows),
                "deleted_old_rows": deleted,
            }
        )

    oct = wb[SHEETS["2026-10"]]
    oct.sheet_view.rightToLeft = True
    r = oct.max_row + 2
    oct.cell(r, 1).value = f"7 | תשלומים ללא חודש מאומת — {len(date_unknown)} רשומות"
    r += 1
    for c, h in enumerate(HEADERS, start=1):
        oct.cell(r, c).value = h
    r += 1
    for pay in date_unknown:
        oct.cell(r, 1).value = pay.get("paymentDate") or "—"
        oct.cell(r, 2).value = KIND_HE.get("expense_date_unknown")
        oct.cell(r, 3).value = pay.get("counterparty")
        oct.cell(r, 4).value = pay.get("description")
        oct.cell(r, 5).value = q2(pay.get("net", 0))
        oct.cell(r, 6).value = q2(pay.get("vat", 0))
        oct.cell(r, 7).value = q2(pay.get("gross", 0))
        oct.cell(r, 8).value = pay.get("allocation")
        oct.cell(r, 9).value = pay.get("sourceId")
        oct.cell(r, 10).value = pay.get("note")
        r += 1

    r += 1
    oct.cell(r, 1).value = f"8 | התחייבויות / יתרות לתשלום — {len(commitments)} רשומות (לא בוצעו בחודש)"
    r += 1
    for c, h in enumerate(HEADERS[:8], start=1):
        oct.cell(r, c).value = h
    r += 1
    for pay in commitments[:40]:
        oct.cell(r, 3).value = pay.get("counterparty")
        oct.cell(r, 4).value = pay.get("description")
        oct.cell(r, 5).value = q2(pay.get("net", 0))
        oct.cell(r, 6).value = q2(pay.get("vat", 0))
        oct.cell(r, 7).value = q2(pay.get("gross", 0))
        oct.cell(r, 8).value = pay.get("allocation")
        oct.cell(r, 9).value = pay.get("sourceId")
        r += 1

    wb.save(XLSX)

    excel_open = None
    try:
        import win32com.client  # type: ignore

        excel = win32com.client.Dispatch("Excel.Application")
        excel.Visible = True
        excel.DisplayAlerts = False
        excel.Workbooks.Open(str(XLSX.resolve()))
        excel_open = "opened"
    except Exception as exc:  # noqa: BLE001
        excel_open = str(exc)

    report = {"sheets": stats, "date_unknown_count": len(date_unknown), "excel_com": excel_open}
    REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False))


if __name__ == "__main__":
    main()
