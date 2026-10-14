import json
import re
from pathlib import Path
from openpyxl import load_workbook

p = Path(r"c:\Users\ERAN YOSEF\Desktop\final projects\FINAL-WEB\projectflow\ProjectFlow_Financial_2026.xlsx")
wb = load_workbook(p, data_only=False)
formulas = []
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                formulas.append({"sheet": ws.title, "cell": cell.coordinate, "f": cell.value[:80]})
Path(r"c:\Users\ERAN YOSEF\Desktop\final projects\FINAL-WEB\projectflow\.tools\excel-formulas.json").write_text(
    json.dumps(formulas, ensure_ascii=False, indent=2), encoding="utf-8"
)
print(len(formulas))
