import json
import re
from pathlib import Path
from openpyxl import load_workbook

XLSX = Path(r"c:\Users\ERAN YOSEF\Desktop\final projects\FINAL-WEB\projectflow\ProjectFlow_Financial_2026.xlsx")
SHEET_YM = {
    "01 ינואר": "2026-01",
    "02 פברואר": "2026-02",
    "03 מרץ": "2026-03",
    "04 אפריל": "2026-04",
    "05 מאי": "2026-05",
    "06 יוני": "2026-06",
    "07 יולי": "2026-07",
    "08 אוגוסט": "2026-08",
    "09 ספטמבר": "2026-09",
    "10 אוקטובר": "2026-10",
}
OVERVIEW_ROWS = [
    ("2026-01", 5),
    ("2026-02", 6),
    ("2026-03", 7),
    ("2026-04", 8),
    ("2026-05", 9),
    ("2026-06", 10),
    ("2026-07", 11),
    ("2026-08", 12),
    ("2026-09", 13),
    ("2026-10", 14),
]
UUID_RE = re.compile(r"^[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}$", re.I)

wb = load_workbook(XLSX)
counts = {}
for ws in wb.worksheets:
    if ws.title not in SHEET_YM:
        continue
    ym = SHEET_YM[ws.title]
    c = 0
    for row in range(1, ws.max_row + 1):
        if ws.cell(row, 2).value != "הוצאה":
            continue
        v = ws.cell(row, 13).value
        if v and UUID_RE.match(str(v).strip()):
            c += 1
    counts[ym] = c
ov = wb["סקירה שנתית"]
for ym, r in OVERVIEW_ROWS:
    ov.cell(r, 5).value = counts[ym]
wb.save(XLSX)
print(json.dumps(counts, indent=2))
