import json
from pathlib import Path
from openpyxl import load_workbook

p = Path(r"c:\Users\ERAN YOSEF\Desktop\final projects\FINAL-WEB\projectflow\ProjectFlow_Financial_2026.xlsx")
wb = load_workbook(p, data_only=True)
out = {"sheets": wb.sheetnames, "month_expense_rows": {}}
for ws in wb.worksheets:
    title = ws.title
    if not any(x in title for x in ("01", "02", "ינואר", "פברואר", "1 ", "2 ")):
        continue
    rows = []
    for r in ws.iter_rows(values_only=True):
        if not r or not any(x is not None for x in r):
            continue
        line = [x for x in r]
        s = " | ".join(str(x) for x in line if x is not None)
        if any(k in s for k in ("800", "28050", "33099", "תותח", "01/01", "27/02", "28/02", "2026")):
            rows.append(line)
    out["month_expense_rows"][title] = rows
Path(r"c:\Users\ERAN YOSEF\Desktop\final projects\FINAL-WEB\projectflow\.tools\excel-dup-snip.json").write_text(
    json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8"
)
