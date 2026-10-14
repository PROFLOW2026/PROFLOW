"""
Remove verified void expenses from ProjectFlow_Financial_2026.xlsx (in-place).
READ-ONLY Production verification must be done separately before running.
"""
from __future__ import annotations

import json
import re
import uuid
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(r"c:\Users\ERAN YOSEF\Desktop\final projects\FINAL-WEB\projectflow")
XLSX = ROOT / "ProjectFlow_Financial_2026.xlsx"
REPORT = ROOT / ".tools" / "excel-remove-voids-report.json"

VOID_IDS = {
    "5642f1de-b5f7-474d-888a-8686b88b2167",
    "7d127420-d37a-469a-97b5-702389de4fee",
    "2807943d-8fcb-41e5-989d-3ee46982f8b0",
    "c16d5480-8422-43b4-9a24-68bc947e8314",
    "e67a342d-bf56-4a79-9b2f-6a77689cf9a5",
    "b59035cf-fe06-49a6-9cb3-97459bed9b4f",
    "989cc273-cd8e-4932-b205-6869d5154b02",
}

# Production finalized expense counts / recognized NET (expense_date month)
PROD = {
    "2026-01": {"cnt": 11, "net": Decimal("54024.745762")},
    "2026-02": {"cnt": 13, "net": Decimal("38621.864406")},
    "2026-03": {"cnt": 11, "net": Decimal("58687.118644")},
    "2026-04": {"cnt": 12, "net": Decimal("45742.203389")},
    "2026-05": {"cnt": 11, "net": Decimal("46275.254237")},
    "2026-06": {"cnt": 12, "net": Decimal("50351.525423")},
    "2026-07": {"cnt": 14, "net": Decimal("71114.525424")},
    "2026-08": {"cnt": 13, "net": Decimal("74425.254235")},
    "2026-09": {"cnt": 11, "net": Decimal("63829.491524")},
    "2026-10": {"cnt": 5, "net": Decimal("2849.830508")},
}

OVERVIEW_ROWS: list[tuple[str, int]] = [
    ("2026-01", 5),
    ("2026-02", 6),
    ("2026-03", 7),
    ("2026-04", 8),
    ("2026-05", 9),
    ("2026-06", 10),
    ("2026-07", 11),
    ("2026-08", 12),
    ("2026-09", 13),
    ("2026-10", 14),
]

SHEET_YM = {
    "01 ינואר": "2026-01",
    "02 פברואר": "2026-02",
    "03 מרץ": "2026-03",
    "04 אפריל": "2026-04",
    "05 מאי": "2026-05",
    "06 יוני": "2026-06",
    "07 יולי": "2026-07",
    "08 אוגוסט": "2026-08",
    "09 ספטמבר": "2026-09",
    "10 אוקטובר": "2026-10",
}

UUID_RE = re.compile(
    r"^[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}$",
    re.I,
)


def q2(v: Decimal) -> Decimal:
    return v.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def as_decimal(v) -> Decimal:
    if v is None or v == "":
        return Decimal("0")
    return Decimal(str(v))


def cell_id(ws, row: int) -> str | None:
    v = ws.cell(row, 13).value
    if v is None:
        return None
    s = str(v).strip()
    if UUID_RE.match(s):
        return s.lower()
    return None


def is_expense_row(ws, row: int) -> bool:
    t = ws.cell(row, 2).value
    return t == "הוצאה" and cell_id(ws, row) is not None


def main() -> None:
    lock = ROOT / "~$ProjectFlow_Financial_2026.xlsx"
    if lock.exists():
        raise SystemExit(
            "Excel lock file exists (~$...). Close ProjectFlow_Financial_2026.xlsx and retry."
        )

    wb = load_workbook(XLSX)
    removed: list[dict] = []

    for ws in wb.worksheets:
        ws.sheet_view.rightToLeft = True
        if ws.title not in SHEET_YM:
            continue
        to_delete: list[int] = []
        for row in range(1, ws.max_row + 1):
            eid = cell_id(ws, row)
            if eid and eid in VOID_IDS and is_expense_row(ws, row):
                to_delete.append(row)
                removed.append(
                    {
                        "sheet": ws.title,
                        "row": row,
                        "id": eid,
                        "net": ws.cell(row, 5).value,
                        "recognized": ws.cell(row, 12).value,
                        "balance": ws.cell(row, 9).value,
                    }
                )
        for row in sorted(to_delete, reverse=True):
            ws.delete_rows(row, 1)

        # Section header row 4
        hdr = ws.cell(4, 1).value
        if hdr and str(hdr).startswith("1 | הוצאות"):
            ym = SHEET_YM[ws.title]
            n = PROD[ym]["cnt"]
            ws.cell(4, 1).value = f"1 | הוצאות — {n} רשומות (סופיות / תקפות)"

        sub = ws.cell(2, 1).value
        if sub and "ללא שינוי נתונים" in str(sub):
            ws.cell(2, 1).value = str(sub).replace(
                "ללא שינוי נתונים",
                "הוצאות מבוטלות (void) הוסרו מתצוגה — נתוני Production",
            )

    # Annual overview counts (col 5) + expense NET col 3
    overview = wb["סקירה שנתית"]
    overview.sheet_view.rightToLeft = True
    note = overview.cell(2, 1).value
    if note and "11.10.2026" in str(note):
        overview.cell(2, 1).value = (
            str(note)
            + " | הוצאות מבוטלות הוסרו מגיליונות החודש (7 שורות)."
        )
    for ym, r in OVERVIEW_ROWS:
        overview.cell(r, 5).value = PROD[ym]["cnt"]
        overview.cell(r, 3).value = float(q2(PROD[ym]["net"]))

    wb["הסבר לקריאה"].sheet_view.rightToLeft = True

    if len(removed) != 7:
        raise SystemExit(f"Expected to remove 7 rows, found {len(removed)}: {removed}")

    wb.save(XLSX)

    # Post-save validation (values)
    wb2 = load_workbook(XLSX, data_only=True)
    hits = []
    expense_count = 0
    monthly_rec: dict[str, Decimal] = {ym: Decimal("0") for ym in PROD}
    for ws in wb2.worksheets:
        if ws.title not in SHEET_YM:
            continue
        ym = SHEET_YM[ws.title]
        for row in range(1, ws.max_row + 1):
            eid = cell_id(ws, row)
            if eid:
                if eid in VOID_IDS:
                    hits.append({"sheet": ws.title, "row": row, "id": eid})
            if is_expense_row(ws, row):
                expense_count += 1
                monthly_rec[ym] += as_decimal(ws.cell(row, 12).value)

    mismatches = []
    for ym, spec in PROD.items():
        excel_cnt = sum(
            1
            for ws in wb2.worksheets
            if ws.title in SHEET_YM and SHEET_YM[ws.title] == ym
            for row in range(1, ws.max_row + 1)
            if is_expense_row(ws, row)
        )
        if excel_cnt != spec["cnt"]:
            mismatches.append({"ym": ym, "excel_cnt": excel_cnt, "prod_cnt": spec["cnt"]})
        delta = q2(monthly_rec[ym]) - q2(spec["net"])
        if abs(delta) > Decimal("0.02"):
            mismatches.append(
                {
                    "ym": ym,
                    "excel_recognized": str(q2(monthly_rec[ym])),
                    "prod_net": str(q2(spec["net"])),
                    "delta": str(delta),
                }
            )

    report = {
        "removed_rows": removed,
        "void_hits_remaining": hits,
        "expense_rows_total": expense_count,
        "mismatches": mismatches,
        "rtl_all_sheets": all(getattr(ws.sheet_view, "rightToLeft", False) for ws in wb2.worksheets),
    }
    REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")

    if hits:
        raise SystemExit(f"Void IDs still present in workbook: {hits}")
    if expense_count != 113:
        raise SystemExit(f"Expected 113 expense rows, got {expense_count}")
    if mismatches:
        raise SystemExit(f"Production mismatch: {mismatches}")

    # Try Excel COM open (Windows)
    open_ok = None
    try:
        import win32com.client  # type: ignore

        excel = win32com.client.Dispatch("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        wb_com = excel.Workbooks.Open(str(XLSX.resolve()))
        open_ok = wb_com.Name == XLSX.name
        wb_com.Close(SaveChanges=False)
        excel.Quit()
    except Exception as exc:  # noqa: BLE001
        open_ok = f"skipped: {exc}"

    report["excel_com_open"] = open_ok
    REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"ok": True, "removed": 7, "expenses": 113, "excel_com_open": open_ok}, indent=2))


if __name__ == "__main__":
    main()
