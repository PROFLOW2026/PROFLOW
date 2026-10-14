import json
from pathlib import Path
from openpyxl import load_workbook

p = Path(r"c:\Users\ERAN YOSEF\Desktop\final projects\FINAL-WEB\projectflow\ProjectFlow_Financial_2026.xlsx")
wb = load_workbook(p, data_only=True)
ws = wb["01 ינואר"]
rows = []
for r in range(1, ws.max_row + 1):
    row = [ws.cell(r, c).value for c in range(1, 14)]
    if any(x is not None for x in row):
        rows.append({"r": r, "row": row})
Path(r"c:\Users\ERAN YOSEF\Desktop\final projects\FINAL-WEB\projectflow\.tools\excel-jan-all.json").write_text(
    json.dumps(rows, ensure_ascii=False, indent=2, default=str), encoding="utf-8"
)
