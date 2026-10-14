"""
BMS (manual lump) + engine NET alignment for payment-based closeout.
Edits ProjectFlow_Financial_2026.xlsx in place; opens via Excel COM when available.
"""
from __future__ import annotations

import json
import re
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(r"c:\Users\ERAN YOSEF\Desktop\final projects\FINAL-WEB\projectflow")
XLSX = ROOT / "ProjectFlow_Financial_2026.xlsx"
REPORT = ROOT / ".tools" / "excel-bms-closeout-report.json"

BMS_ID = "bbd3cdce-22e5-477f-bae1-1161e87cacb6"
BMS_GROSS = Decimal("8689")
BMS_NET = Decimal("7363.559322")
BMS_VAT = BMS_GROSS - BMS_NET

SHEET_YM = {
    "01 ינואר": "2026-01",
    "02 פברואר": "2026-02",
    "03 מרץ": "2026-03",
    "04 אפריל": "2026-04",
    "05 מאי": "2026-05",
    "06 יוני": "2026-06",
    "07 יולי": "2026-07",
    "08 אוגוסט": "2026-08",
    "09 ספטמבר": "2026-09",
    "10 אוקטובר": "2026-10",
}

# Payment-based engine NET (post BMS date-unknown policy)
ENGINE_NET = {
    "2026-01": Decimal("47133.138474"),
    "2026-02": Decimal("40534.494406"),
    "2026-03": Decimal("53062.111695"),
    "2026-04": Decimal("48628.213389"),
    "2026-05": Decimal("48486.264237"),
    "2026-06": Decimal("52562.555423"),
    "2026-07": Decimal("73085.425424"),
    "2026-08": Decimal("67925.027795"),
    "2026-09": Decimal("57377.941185"),
    "2026-10": Decimal("3761.840508"),
}

OVERVIEW_ROWS = [
    ("2026-01", 5),
    ("2026-02", 6),
    ("2026-03", 7),
    ("2026-04", 8),
    ("2026-05", 9),
    ("2026-06", 10),
    ("2026-07", 11),
    ("2026-08", 12),
    ("2026-09", 13),
    ("2026-10", 14),
]

UUID_RE = re.compile(
    r"^[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}$",
    re.I,
)


def q2(v: Decimal) -> float:
    return float(v.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def cell_id(ws, row: int) -> str | None:
    v = ws.cell(row, 13).value
    if not v:
        return None
    s = str(v).strip().lower()
    return s if UUID_RE.match(s) else None


def find_row_by_prefix(ws, prefix: str) -> int | None:
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if v and str(v).startswith(prefix):
            return r
    return None


def main() -> None:
    lock = ROOT / "~$ProjectFlow_Financial_2026.xlsx"
    if lock.exists():
        raise SystemExit("Close Excel (lock file) and retry.")

    wb = load_workbook(XLSX)
    cleared: list[dict] = []

    for sheet, ym in SHEET_YM.items():
        ws = wb[sheet]
        ws.sheet_view.rightToLeft = True
        for row in range(1, ws.max_row + 1):
            if cell_id(ws, row) != BMS_ID:
                continue
            old = ws.cell(row, 12).value
            if old not in (None, "", 0):
                cleared.append({"sheet": sheet, "row": row, "was": old})
            ws.cell(row, 12).value = None
            status = ws.cell(row, 11).value or ""
            tag = "שולם — תאריך תשלום לא מאומת (חריגה)"
            if tag not in str(status):
                ws.cell(row, 11).value = f"{status} | {tag}".strip(" |")
            ws.cell(row, 8).value = q2(BMS_GROSS)
            ws.cell(row, 9).value = 0

    oct = wb["10 אוקטובר"]
    oct.sheet_view.rightToLeft = True
    hdr = find_row_by_prefix(oct, "6 |")
    insert_at = hdr if hdr else oct.max_row + 2
    if not hdr:
        oct.cell(insert_at, 1).value = "6 | חריגות — תשלום ללא חודש מאומת — 1 רשומות"
        insert_at += 1
        headers = [
            "תאריך רישום",
            "סוג",
            "ספק",
            "תיאור",
            "NET ₪",
            "מע״מ ₪",
            "GROSS ₪",
            "שולם GROSS",
            "יתרה",
            "הערה",
            "סטטוס",
            "מוכר בעלות",
            "מזהה",
        ]
        for c, h in enumerate(headers, start=1):
            oct.cell(insert_at, c).value = h
        insert_at += 1
    data_row = insert_at
    oct.cell(data_row, 1).value = "2026-09-06"
    oct.cell(data_row, 2).value = "חריגת תשלום"
    oct.cell(data_row, 3).value = "ב.מ.ש"
    oct.cell(data_row, 4).value = "טיפולי רכב — אישור ידני מרוכז (12 פעימות)"
    oct.cell(data_row, 5).value = q2(BMS_NET)
    oct.cell(data_row, 6).value = q2(BMS_VAT)
    oct.cell(data_row, 7).value = q2(BMS_GROSS)
    oct.cell(data_row, 8).value = q2(BMS_GROSS)
    oct.cell(data_row, 9).value = 0
    oct.cell(data_row, 10).value = "audit 06/09/2026 — לא הוכח חודש תזרים בפועל"
    oct.cell(data_row, 11).value = "PAID_DATE_UNKNOWN"
    oct.cell(data_row, 12).value = 0
    oct.cell(data_row, 13).value = BMS_ID

    overview = wb["סקירה שנתית"]
    overview.sheet_view.rightToLeft = True
    for ym, r in OVERVIEW_ROWS:
        overview.cell(r, 3).value = q2(ENGINE_NET[ym])

    wb.save(XLSX)

    excel_open = None
    try:
        import win32com.client  # type: ignore

        excel = win32com.client.Dispatch("Excel.Application")
        excel.Visible = True
        excel.DisplayAlerts = False
        excel.Workbooks.Open(str(XLSX.resolve()))
        excel_open = "opened"
    except Exception as exc:  # noqa: BLE001
        excel_open = str(exc)

    report = {
        "bms_id": BMS_ID,
        "cleared_recognition_cells": cleared,
        "engine_net_september": str(ENGINE_NET["2026-09"]),
        "excel_com": excel_open,
    }
    REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
